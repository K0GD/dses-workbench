"""Build a DSES-styled PDF from a Markdown source (e.g. Installing.md).

The deliverable is the PDF:
    DSES_Radio_Astronomy_Workbench_Installation.pdf

Internally the document is rendered to a temporary .docx (python-docx, styled
to match the DSES house template — margins, title color, Heading 1
white-on-teal banner, Heading 3 dark-teal, page header with "DSES" + subtitle,
footer with date + page number) and then converted to PDF. The .docx is a
throwaway intermediate — it is written to a temp file and removed after the
PDF is built, unless you pass --docx to keep it. PDF conversion uses
LibreOffice on macOS/Linux and Microsoft Word on Windows (see
convert_docx_to_pdf).

Handles the Markdown subset used in Installing.md:
    # / ## / ###       — headings
    paragraphs         — regular text with inline **bold**, *italic*, `code`
    - / *              — bulleted lists
    1.                 — numbered lists
    ```                — fenced code blocks (monospace, light-gray background)
    | a | b |          — tables (with --- alignment separator)
    [text](url)        — hyperlinks (rendered as text with the URL after, since
                         full hyperlink XML in python-docx is verbose)

Run from the project root with the project's conda env:
    .conda\\python.exe build_install_docx.py
"""

from __future__ import annotations

import html
import os
import re
import sys
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------------------------------------------------------------------------
# Smart-quote pre-pass: convert ASCII apostrophes and double quotes in body
# text to typographer equivalents, leaving fenced code blocks and inline
# `code` spans untouched. Matches DSES house style (Word's AutoCorrect would
# do the same thing for hand-typed text, but python-docx writes text
# verbatim and Word's autocorrect doesn't run on programmatic content).
# ---------------------------------------------------------------------------

LEFT_DOUBLE  = '“'   # "
RIGHT_DOUBLE = '”'   # "
APOSTROPHE   = '’'   # ’ (also doubles as right single quote)


def smartify_quotes(text: str) -> str:
    out = []
    in_fence = False
    dq_state = {'open': True}  # next " is opening
    for line in text.split('\n'):
        if line.lstrip().startswith('```'):
            in_fence = not in_fence
            out.append(line)
            continue
        if in_fence:
            out.append(line)
            continue
        # Reset double-quote state on paragraph breaks so an unmatched quote
        # in one paragraph doesn't flip the polarity for the next.
        if line.strip() == '':
            dq_state['open'] = True
        out.append(_smartify_line(line, dq_state))
    return '\n'.join(out)


def _smartify_line(line: str, dq_state: dict) -> str:
    result = []
    i = 0
    n = len(line)
    while i < n:
        c = line[i]
        if c == '`':
            # Inline code span — emit verbatim through the closing backtick.
            j = line.find('`', i + 1)
            if j == -1:
                result.append(c)
                i += 1
                continue
            result.append(line[i:j + 1])
            i = j + 1
            continue
        if c == "'":
            # Apostrophe only if it's after a letter (contraction / possessive).
            # Leave others straight — they might be opening single quotes,
            # feet markers, or other punctuation.
            prev = line[i - 1] if i > 0 else ''
            if prev.isalpha():
                result.append(APOSTROPHE)
            else:
                result.append(c)
            i += 1
            continue
        if c == '"':
            result.append(LEFT_DOUBLE if dq_state['open'] else RIGHT_DOUBLE)
            dq_state['open'] = not dq_state['open']
            i += 1
            continue
        result.append(c)
        i += 1
    return ''.join(result)


# Defaults for the install guide; CLI options can override. Used as module-
# level constants because the cover-page / header builders read them.
SRC = Path("Installing.md")
DST_PDF  = Path("DSES_Radio_Astronomy_Workbench_Installation.pdf")
DOC_TITLE    = "DSES Radio Astronomy Workbench"
DOC_SUBTITLE = "Installation Guide"
DOC_VERSION  = "v1.1.6"
DOC_AUTHOR   = "Richard M Hambly (K0GD)"
DOC_ORG      = "DSES"

# Typography. Aligned with C:\CNS-Systems\DOCUMENT_STANDARDS.md section 3
# (Rick, 19-Aug-2026) so DSES documents match CNS Systems reports: Minion Pro
# body, Myriad Pro headings/display, Source Code Pro for code. The DSES skin
# (teal H1 banner, colors, sizes) is unchanged — only the faces moved.
#
# IMPORTANT: these are OpenType-PS (CFF) faces. Word's SaveAs-PDF silently
# RASTERIZES them (verified 19-Aug-2026: Minion/Myriad runs came out as images
# with no text layer, while TrueType Source Code Pro embedded fine), so the
# Windows converter below prints through the Adobe PDF printer instead. If you
# change these back to TrueType faces (Calibri/Cambria/Consolas), the plain
# SaveAs path is adequate again.
FONT_BODY = "Minion Pro"
FONT_HEAD = "Myriad Pro"
FONT_MONO = "Source Code Pro"

# DSES house style (pulled from EVE-26 + Pulsar installation reference docs)
TITLE_COLOR_RGB     = RGBColor(0x15, 0x60, 0x82)  # teal/blue
H1_BG_HEX           = "156082"                    # same teal as banner fill
H1_TEXT_RGB         = RGBColor(0xFF, 0xFF, 0xFF)  # white on the banner
H3_COLOR_RGB        = RGBColor(0x0A, 0x2F, 0x40)  # dark teal
SUBTITLE_COLOR_RGB  = RGBColor(0x59, 0x59, 0x59)  # mid-gray
CODE_FILL_HEX       = "F2F2F2"                    # very light gray
PAGE_MARGINS = dict(  # inches; matches both DSES reference documents
    left=1.00, right=0.81, top=1.36, bottom=1.00,
)
# Optional logo shown right-justified in the page header (pages 2+). Set via
# --header-logo; None keeps the classic text-only header.
HEADER_LOGO = None


# ---------------------------------------------------------------------------
# Inline run handling — split "...some **bold** and `code` text..." into a
# list of (text, style_dict) tuples.
# ---------------------------------------------------------------------------

INLINE_RE = re.compile(
    r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))'
)
LINK_RE = re.compile(r'\[([^\]]+)\]\(([^)]+)\)')


def parse_inline(text: str):
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            yield text[pos:m.start()], {}
        token = m.group(0)
        if token.startswith('**'):
            yield token[2:-2], {'bold': True}
        elif token.startswith('`'):
            yield token[1:-1], {'code': True}
        elif token.startswith('['):
            lm = LINK_RE.match(token)
            if lm:
                label, url = lm.group(1), lm.group(2)
                if label.strip() == url.strip():
                    yield label, {'code': True}
                else:
                    yield label, {}
                    yield f" ({url})", {'code': True}
        elif token.startswith('*'):
            yield token[1:-1], {'italic': True}
        pos = m.end()
    if pos < len(text):
        yield text[pos:], {}


def add_runs(paragraph, text: str):
    for chunk, style in parse_inline(text):
        if not chunk:
            continue
        chunk = html.unescape(chunk)
        run = paragraph.add_run(chunk)
        if style.get('bold'):
            run.bold = True
        if style.get('italic'):
            run.italic = True
        if style.get('code'):
            run.font.name = FONT_MONO
            run.font.size = Pt(9.5)


# ---------------------------------------------------------------------------
# DSES style primitives — small wrappers around the raw OOXML so the body
# of the document builder stays readable.
# ---------------------------------------------------------------------------

def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color_hex)
    shd.set(qn('w:val'), 'clear')
    tc_pr.append(shd)


def set_paragraph_shading(paragraph, color_hex):
    """Paint a solid fill behind a paragraph (used for the DSES H1 banner)."""
    pPr = paragraph._element.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    pPr.append(shd)


def add_page_number_field(paragraph):
    """Insert a PAGE field code so Word numbers pages live."""
    run = paragraph.add_run()
    for tag in ('begin', None, 'end'):
        if tag is None:
            instr = OxmlElement('w:instrText')
            instr.text = 'PAGE'
            run._r.append(instr)
        else:
            ch = OxmlElement('w:fldChar')
            ch.set(qn('w:fldCharType'), tag)
            run._r.append(ch)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Cover page + section header/footer (DSES house style)
# ---------------------------------------------------------------------------

def configure_section(section):
    section.left_margin   = Inches(PAGE_MARGINS['left'])
    section.right_margin  = Inches(PAGE_MARGINS['right'])
    section.top_margin    = Inches(PAGE_MARGINS['top'])
    section.bottom_margin = Inches(PAGE_MARGINS['bottom'])

    # Suppress the header/footer on the cover page: enable a distinct
    # first-page header/footer and leave it empty. The primary header/footer
    # configured below then applies only to page 2 onward.
    section.different_first_page_header_footer = True

    # Header (pages 2+): document title on line 1, subtitle on line 2.
    hdr = section.header
    # Reuse existing first paragraph; we control its content fully.
    p1 = hdr.paragraphs[0]
    p1.text = ''
    r = p1.add_run(DOC_TITLE)
    r.bold = True
    r.font.name = FONT_HEAD
    r.font.size = Pt(11)
    r.font.color.rgb = TITLE_COLOR_RGB
    # Optional logo, right-justified on the title line via a right tab stop
    # at the text-column edge.
    if HEADER_LOGO and Path(HEADER_LOGO).is_file():
        content_w = 8.5 - PAGE_MARGINS['left'] - PAGE_MARGINS['right']
        p1.paragraph_format.tab_stops.add_tab_stop(
            Inches(content_w), WD_TAB_ALIGNMENT.RIGHT)
        # The built-in Header style carries center (3.25") and right (6.5")
        # stops; a single tab would land on the center one and park the logo
        # mid-page. Emit w:val="clear" entries so only our stop remains.
        # NOTE: w:tab elements must be in ascending w:pos order or Word
        # discards the list — insert 9360 first, then 4680 in front of it.
        tabs_el = p1._p.pPr.find(qn('w:tabs'))
        for pos_twips in ('9360', '4680'):          # 6.5" then 3.25"
            clear = OxmlElement('w:tab')
            clear.set(qn('w:val'), 'clear')
            clear.set(qn('w:pos'), pos_twips)
            tabs_el.insert(0, clear)
        p1.add_run('\t')
        p1.add_run().add_picture(str(HEADER_LOGO), height=Inches(0.42))
    p2 = hdr.add_paragraph()
    r2 = p2.add_run(DOC_SUBTITLE)
    r2.italic = True
    r2.font.name = FONT_HEAD
    r2.font.size = Pt(9)
    r2.font.color.rgb = SUBTITLE_COLOR_RGB

    # Footer: date <tab> page number — matches the reference DSES docs.
    fp = section.footer.paragraphs[0]
    fp.text = ''
    today = date.today().strftime("%d-%b-%y")  # e.g. "16-May-26"
    fr = fp.add_run(today)
    fr.font.size = Pt(9)
    fr.font.color.rgb = SUBTITLE_COLOR_RGB
    fp.add_run('\t')
    add_page_number_field(fp)
    for run in fp.runs:
        run.font.name = FONT_HEAD
        run.font.size = Pt(9)
        run.font.color.rgb = SUBTITLE_COLOR_RGB


def _toc_plain(title):
    """Strip the small subset of inline markdown the headings use, for a clean
    table-of-contents entry."""
    t = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', title)   # [text](url) -> text
    return t.replace('**', '').replace('`', '').strip()


def collect_headings(lines):
    """(level, title) for every ATX heading in the markdown, skipping fenced
    code blocks so a `#` comment inside a code sample isn't mistaken for one."""
    out = []
    in_code = False
    for ln in lines:
        s = ln.strip()
        if s.startswith('```'):
            in_code = not in_code
            continue
        if in_code:
            continue
        if s.startswith('### '):
            out.append((3, _toc_plain(s[4:])))
        elif s.startswith('## '):
            out.append((2, _toc_plain(s[3:])))
        elif s.startswith('# '):
            out.append((1, _toc_plain(s[2:])))
    return out


def add_table_of_contents(doc, headings):
    """A static, pre-populated outline of the document's headings.

    Used when toc_mode == 'static' (the default off Windows). The Windows
    default is add_toc_field(), a real TOC field that Word evaluates.
    Rationale for keeping this path: Word's TOC *field* renders as placeholder
    text until an application re-evaluates it, and on macOS/Linux the
    LibreOffice conversion path can't update it (and LibreOffice's embedded
    Python is launch-constraint-blocked from scripting it). A static outline
    renders identically and correctly in Word, LibreOffice, and any PDF viewer.
    Trade-off: no live page numbers — the section numbers carried in the
    heading text provide the structure instead."""
    p = doc.add_heading("Contents", level=1)
    set_paragraph_shading(p, H1_BG_HEX)

    # Drop a leading level-1 entry (the document title, already on the cover).
    entries = headings[1:] if (headings and headings[0][0] == 1) else headings
    base = min((lv for lv, _ in entries), default=1)
    for level, title in entries:
        para = doc.add_paragraph()
        para.paragraph_format.left_indent = Pt(18 * (level - base))
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(title)
        run.font.name = FONT_HEAD
        run.font.size = Pt(11)
        if level == base:
            run.font.bold = True

    add_page_break(doc)


def add_banner_paragraph(doc, text):
    """A paragraph with the Heading-1 look (white on the teal banner) but WITHOUT
    the Heading 1 style, so a Word TOC field does not list it. Used for the
    document title in the body and for the "Contents" heading itself."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT_HEAD
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = H1_TEXT_RGB
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    set_paragraph_shading(p, H1_BG_HEX)
    return p


def _ensure_toc_styles(doc):
    """Define Word's built-in 'toc 1..3' paragraph styles in the house look:
    level 1 Myriad bold, lower levels regular, right tab with dot leader at
    the text width so the page numbers line up (DOCUMENT_STANDARDS TOC rule).
    Word applies these when it populates the TOC field. The Hyperlink
    character style is pinned to the house font (the \\h switch makes the
    entries hyperlinks) so no theme font leaks into the PDF."""
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_TAB_LEADER
    text_width = Inches(8.5 - PAGE_MARGINS['left'] - PAGE_MARGINS['right'])
    for lvl in (1, 2, 3):
        name = f'toc {lvl}'
        try:
            st = doc.styles[name]
        except KeyError:
            st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        # Word recognises its built-in TOC styles by style ID 'TOC1'..'TOC3'
        # (python-docx derives 'toc1'); with the wrong ID Word treats ours as
        # custom, renames it 'TOC 11', and formats the entries with its own
        # defaults (theme fonts leaked into the PDF, 2026-09-10).
        st.element.set(qn('w:styleId'), f'TOC{lvl}')
        # python-docx marks added styles customStyle=1; a custom style with a
        # built-in name is exactly what makes Word rename it. Drop the flag.
        st.element.attrib.pop(qn('w:customStyle'), None)
        st.base_style = doc.styles['Normal']
        st.font.name = FONT_HEAD
        st.font.size = Pt(10.5 if lvl == 1 else 10)
        st.font.bold = (lvl == 1)
        st.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
        pf = st.paragraph_format
        pf.left_indent = Pt(14 * (lvl - 1))
        pf.space_before = Pt(5 if lvl == 1 else 0)
        pf.space_after = Pt(2)
        pf.tab_stops.clear_all()
        pf.tab_stops.add_tab_stop(text_width, WD_TAB_ALIGNMENT.RIGHT,
                                  WD_TAB_LEADER.DOTS)
    try:
        hl = doc.styles['Hyperlink']
    except KeyError:
        hl = doc.styles.add_style('Hyperlink', WD_STYLE_TYPE.CHARACTER)
        hl.element.attrib.pop(qn('w:customStyle'), None)
    hl.font.name = FONT_HEAD
    hl.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    hl.font.underline = False


def add_toc_field(doc):
    """Insert a real Word TOC field (levels 1-3, hyperlinked, page numbers).
    It renders as a placeholder line until an application evaluates it:
    the Windows/Word conversion path does that (Fields.Update twice), so the
    PDF carries live page numbers. LibreOffice does NOT evaluate it, which
    is why the static outline (add_table_of_contents) remains the default
    off Windows."""
    add_banner_paragraph(doc, "Contents")
    _ensure_toc_styles(doc)
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    begin.set(qn('w:dirty'), 'true')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = ' TOC \\o "1-3" \\h \\z '
    sep = OxmlElement('w:fldChar')
    sep.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t')
    txt.text = "Table of contents: open in Word and update fields (F9)."
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for el in (begin, instr, sep, txt, end):
        run._r.append(el)
    add_page_break(doc)


def add_cover_page(doc):
    # Two empty paragraphs at the top push the title down a few cm
    for _ in range(2):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run(DOC_TITLE)
    tr.font.name = FONT_HEAD
    tr.font.size = Pt(38)
    tr.font.color.rgb = TITLE_COLOR_RGB
    tr.bold = True

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run(DOC_SUBTITLE)
    sr.font.name = FONT_HEAD
    sr.font.size = Pt(24)
    sr.font.color.rgb = SUBTITLE_COLOR_RGB
    sr.italic = True

    for _ in range(2):
        doc.add_paragraph()

    # If the project ships a PNG icon, drop it on the cover as a small mark.
    icon_path = Path('icons/dses_workbench.png')
    if icon_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(icon_path), width=Inches(1.5))

    for _ in range(3):
        doc.add_paragraph()

    # Version / author / org / date stacked at the bottom of the cover.
    # DOC_VERSION starting with 'v' renders as "Version X.Y.Z" (manuals);
    # anything else (e.g. "Rev A" for reports) is printed verbatim.
    ver_text = (f"Version {DOC_VERSION.lstrip('v')}"
                if DOC_VERSION.startswith('v') else DOC_VERSION)
    for text in (
        ver_text,
        DOC_AUTHOR,
        DOC_ORG,
        date.today().strftime("%B %Y"),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name = FONT_HEAD
        r.font.size = Pt(16)
        r.font.color.rgb = SUBTITLE_COLOR_RGB
    add_page_break(doc)


def apply_heading_styles(doc):
    """Tweak the built-in heading styles to match DSES (after the doc is
    created, since python-docx populates styles from the default template)."""
    h1 = doc.styles['Heading 1']
    h1.font.name = FONT_HEAD
    h1.font.color.rgb = H1_TEXT_RGB
    h1.font.size = Pt(11)
    h1.font.bold = True
    # Rick's H1 spacing tweak: 12 pt above, 6 pt below (was 24/0).
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    # Heading-2 stays default (slate); Heading 3 → dark teal
    h2 = doc.styles['Heading 2']
    h2.font.name = FONT_HEAD
    h3 = doc.styles['Heading 3']
    h3.font.name = FONT_HEAD
    h3.font.color.rgb = H3_COLOR_RGB
    h3.font.bold = True

    normal = doc.styles['Normal']
    normal.font.name = FONT_BODY
    normal.font.size = Pt(11)

    # The default template's heading styles name the THEME fonts
    # (asciiTheme="majorHAnsi" = Cambria) and theme attributes take precedence
    # over the explicit face python-docx's font.name sets. Word then copies
    # the heading paragraph-mark formatting into the TOC entries, so Cambria
    # leaked into the PDF's embedded fonts (2026-09-10). Strip the theme
    # attributes and pin every script slot to the house heading face.
    for st in (h1, h2, h3):
        rpr = st.element.get_or_add_rPr()
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.insert(0, rf)
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            rf.attrib.pop(qn(f'w:{attr}'), None)
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            rf.set(qn(f'w:{attr}'), FONT_HEAD)

    # Also pin the DOCUMENT DEFAULTS (w:docDefaults), not just the Normal
    # style: runless paragraphs (spacers, image anchors, table paragraph
    # marks) fall back to Word's built-in default (Times New Roman / theme
    # Calibri), which then leaks into the PDF's embedded-font list. Same
    # rule as DOCUMENT_STANDARDS.md's generator note, ported to python-docx.
    styles_el = doc.styles.element
    dd = styles_el.find(qn('w:docDefaults'))
    if dd is None:
        dd = OxmlElement('w:docDefaults')
        styles_el.insert(0, dd)
    rprd = dd.find(qn('w:rPrDefault'))
    if rprd is None:
        rprd = OxmlElement('w:rPrDefault')
        dd.insert(0, rprd)
    rpr = rprd.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        rprd.append(rpr)
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(attr), FONT_BODY)


# ---------------------------------------------------------------------------
# Code blocks + tables (unchanged from previous version)
# ---------------------------------------------------------------------------

def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, CODE_FILL_HEX)
    cell.text = ''
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = p.add_run(line)
        run.font.name = FONT_MONO
        run.font.size = Pt(9.5)


def add_table(doc, header_row, body_rows, widths_in=None):
    """Markdown table -> Word table. Rows are marked cantSplit (a row never
    breaks across a page) and the header row repeats at the top of every
    page the table spans (Rick, 2026-09-10). `widths_in` (list of inches,
    one per column) fixes the column widths; without it Word autofits."""
    cols = len(header_row)
    table = doc.add_table(rows=1 + len(body_rows), cols=cols)
    table.style = 'Light Grid Accent 1'
    if widths_in:
        table.autofit = False
        for j, w in enumerate(widths_in[:cols]):
            table.columns[j].width = Inches(w)
            for row in table.rows:
                row.cells[j].width = Inches(w)
    for r_idx, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit')
        tr_pr.append(cant)
        if r_idx == 0:
            hdr = OxmlElement('w:tblHeader')
            tr_pr.append(hdr)
    for j, cell_text in enumerate(header_row):
        cell = table.rows[0].cells[j]
        cell.text = ''
        add_runs(cell.paragraphs[0], cell_text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.name = FONT_HEAD
    for i, row in enumerate(body_rows):
        for j in range(cols):
            cell_text = row[j] if j < len(row) else ''
            cell = table.rows[1 + i].cells[j]
            cell.text = ''
            add_runs(cell.paragraphs[0], cell_text)


def parse_table_block(lines, start):
    def split_row(s):
        s = s.strip()
        if s.startswith('|'):
            s = s[1:]
        if s.endswith('|'):
            s = s[:-1]
        return [c.strip() for c in s.split('|')]
    i = start
    header = split_row(lines[i]); i += 1
    i += 1  # alignment row
    body = []
    while i < len(lines) and lines[i].strip().startswith('|'):
        body.append(split_row(lines[i]))
        i += 1
    return header, body, i


# ---------------------------------------------------------------------------
# Main parser
# ---------------------------------------------------------------------------

# TOC rendering: 'field' = a real Word TOC field with page numbers (needs the
# Word conversion path to evaluate it); 'static' = pre-populated outline that
# renders anywhere (LibreOffice cannot evaluate the field). Default by platform.
TOC_MODE_DEFAULT = 'field' if sys.platform == 'win32' else 'static'


def md_to_docx(src_path: Path, dst_path: Path, cover: bool = True,
               toc: bool = True, toc_mode: str = None):
    toc_mode = toc_mode or TOC_MODE_DEFAULT
    text = src_path.read_text(encoding='utf-8')
    text = smartify_quotes(text)
    lines = text.split('\n')

    doc = Document()
    configure_section(doc.sections[0])
    apply_heading_styles(doc)
    if cover:
        add_cover_page(doc)
    if toc and toc_mode == 'field':
        add_toc_field(doc)
    elif toc:
        add_table_of_contents(doc, collect_headings(lines))

    # Style a Heading-1 paragraph so it gets the teal banner. We do this by
    # post-processing each Heading-1 paragraph after add_heading() rather
    # than wiring it into the global style (which would also paint the TOC,
    # the cover page, etc., with shading).
    # The document's first level-1 heading is its title (already on the
    # cover): render it as a banner paragraph WITHOUT the Heading 1 style so
    # a TOC field does not list it (the static outline drops it likewise).
    seen_h1 = {'n': 0}

    def add_h1_banner(text):
        seen_h1['n'] += 1
        if cover and seen_h1['n'] == 1:
            add_banner_paragraph(doc, text)
            return
        p = doc.add_heading(text, level=1)
        set_paragraph_shading(p, H1_BG_HEX)

    i = 0
    prev_was_table = False          # for post-table paragraph spacing
    pending_widths = None           # from a '<!-- widths: ... -->' comment
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        # A body paragraph immediately under a table gets a little air above
        # it (Rick's house tweak). Latch the "previous block was a table"
        # state here and clear it each iteration; only the table branch re-sets
        # it, so a heading/image between table and paragraph correctly resets.
        after_table = prev_was_table
        prev_was_table = False

        if stripped.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, code_lines)
            continue

        # Single-line HTML comments are authoring notes, except the table
        # width hint '<!-- widths: 1.2,3.0,2.3 -->' (inches per column) that
        # applies to the next table.
        m_c = re.match(r'^<!--\s*(.*?)\s*-->$', stripped)
        if m_c:
            m_w = re.match(r'widths:\s*([\d.,\s]+)$', m_c.group(1))
            if m_w:
                pending_widths = [float(x) for x in m_w.group(1).split(',')
                                  if x.strip()]
            i += 1; continue

        if stripped.startswith('|') and i + 1 < len(lines) \
                and re.match(r'^\s*\|[\s\-:|]+\|\s*$', lines[i + 1]):
            header, body, i = parse_table_block(lines, i)
            add_table(doc, header, body, widths_in=pending_widths)
            pending_widths = None
            prev_was_table = True
            continue

        if stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
            i += 1; continue
        if stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
            i += 1; continue
        if stripped.startswith('# '):
            add_h1_banner(stripped[2:])
            i += 1; continue

        # Image: a line of the form ![caption](path). The path is resolved
        # relative to the source .md; the alt text becomes an italic caption.
        m_img = re.match(r'^!\[(.*?)\]\((.+?)\)\s*$', stripped)
        if m_img:
            caption, img_ref = m_img.group(1), m_img.group(2)
            img_path = Path(img_ref)
            if not img_path.is_absolute():
                img_path = src_path.parent / img_path
            if img_path.is_file():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(str(img_path), width=Inches(6.2))
                if caption:
                    cp = doc.add_paragraph()
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cr = cp.add_run(caption)
                    cr.font.size = Pt(8.5)
                    cr.font.italic = True
                    cr.font.color.rgb = SUBTITLE_COLOR_RGB
            else:
                p = doc.add_paragraph()
                add_runs(p, f"*[missing image: {img_ref}]*")
            i += 1; continue

        if re.match(r'^\s*[-*]\s+', line):
            while i < len(lines) and re.match(r'^\s*[-*]\s+', lines[i]):
                content = re.sub(r'^\s*[-*]\s+', '', lines[i])
                i += 1
                # Coalesce hard-wrapped continuation lines (indented, not a
                # new bullet / number / blank) into the SAME bullet — else
                # they render as detached body paragraphs after the item.
                while (i < len(lines) and lines[i].startswith('  ')
                       and lines[i].strip() != ''
                       and not re.match(r'^\s*[-*]\s+', lines[i])
                       and not re.match(r'^\s*\d+\.\s+', lines[i])):
                    content += ' ' + lines[i].strip()
                    i += 1
                p = doc.add_paragraph(style='List Bullet')
                add_runs(p, content)
            continue

        if re.match(r'^\s*\d+\.\s+', line):
            while i < len(lines) and re.match(r'^\s*\d+\.\s+', lines[i]):
                content = re.sub(r'^\s*\d+\.\s+', '', lines[i])
                p = doc.add_paragraph(style='List Number')
                add_runs(p, content)
                while (i + 1 < len(lines) and lines[i + 1].startswith('   ')
                       and not re.match(r'^\s*\d+\.\s+', lines[i + 1])
                       and not re.match(r'^\s*[-*]\s+', lines[i + 1])
                       and lines[i + 1].strip() != ''):
                    i += 1
                    cont = lines[i].strip()
                    if cont.startswith('```'):
                        code_lines = []
                        i += 1
                        while i < len(lines) and not lines[i].strip().startswith('```'):
                            code_lines.append(lines[i].lstrip())
                            i += 1
                        add_code_block(doc, code_lines)
                    elif cont:
                        cp = doc.add_paragraph()
                        cp.paragraph_format.left_indent = Inches(0.5)
                        add_runs(cp, cont)
                i += 1
            continue

        if stripped == '':
            prev_was_table = after_table   # carry table-state across blank lines
            i += 1; continue

        # Paragraph: coalesce wrapped lines
        para_lines = [line]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            ns = nxt.strip()
            if ns == '' or ns.startswith(('#', '-', '*', '|', '```')) \
                    or re.match(r'^\s*\d+\.\s+', nxt):
                break
            para_lines.append(nxt)
            i += 1
        p = doc.add_paragraph()
        if after_table:
            p.paragraph_format.space_before = Pt(6)
        para_text = ' '.join(s.strip() for s in para_lines)
        add_runs(p, para_text)
        # A lead-in paragraph ("The decisions are:") stays with the list,
        # table, or code block it introduces instead of hanging at the foot
        # of a page (Rick, 2026-09-10).
        if para_text.rstrip().endswith(':'):
            p.paragraph_format.keep_with_next = True

    doc.save(dst_path)
    print(f"Wrote {dst_path} ({dst_path.stat().st_size // 1024} KB)")


def _kill_stale_invisible_word():
    """Kill any orphaned background Word processes left over from a
    previous run that crashed between Documents.Open and word.Quit.
    Skipped silently if psutil isn't available."""
    try:
        import psutil
    except ImportError:
        return
    for p in psutil.process_iter(['name']):
        if (p.info.get('name') or '').lower() == 'winword.exe':
            try:
                p.kill()
            except (psutil.NoSuchProcess, psutil.AccessDenied):
                pass


def _find_soffice():
    """Locate a LibreOffice/OpenOffice headless binary, or None."""
    import shutil
    for name in ("soffice", "libreoffice"):
        p = shutil.which(name)
        if p:
            return p
    # Common macOS install location (not on PATH by default).
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if mac.is_file():
        return str(mac)
    return None


def _convert_docx_to_pdf_soffice(soffice: str, docx_path: Path, pdf_path: Path):
    """macOS/Linux PDF path: convert via LibreOffice headless.

    Note: LibreOffice does not re-evaluate Word's TOC *field* the way Word
    does, so the table of contents may render with its placeholder text. The
    body, headings, page numbers and styling come through fine. For a fully
    populated TOC, build on Windows with Word (the path below) or open the
    .docx in Word once and Save As PDF."""
    import subprocess
    outdir = pdf_path.resolve().parent
    subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                    "--outdir", str(outdir), str(docx_path.resolve())],
                   check=True, capture_output=True, text=True)
    produced = outdir / (docx_path.stem + ".pdf")
    if produced != pdf_path.resolve():
        produced.replace(pdf_path)
    print(f"Wrote {pdf_path} ({pdf_path.stat().st_size // 1024} KB) via LibreOffice")


def convert_docx_to_pdf(docx_path: Path, pdf_path: Path):
    """Convert the DOCX to PDF using the best available engine:

      * Windows + Microsoft Word -> Word via pywin32 (updates the TOC/PAGE
        fields properly; the preferred, fully-correct path);
      * otherwise (macOS/Linux)  -> LibreOffice headless, if installed.

    Raises if neither is available so the caller can fall back to shipping
    the .docx and asking for a manual Save-As-PDF."""
    if sys.platform == "win32":
        return _convert_docx_to_pdf_word(docx_path, pdf_path)

    soffice = _find_soffice()
    if soffice:
        return _convert_docx_to_pdf_soffice(soffice, docx_path, pdf_path)
    raise RuntimeError(
        "No PDF engine found. On Windows install Microsoft Word; on "
        "macOS/Linux install LibreOffice (macOS: `brew install --cask "
        "libreoffice`) and re-run, or open the .docx and Save As PDF.")


def _adobe_pdf_printer_available():
    """True if the 'Adobe PDF' printer (Acrobat Distiller) is installed.

    Note this is independent of Acrobat.exe itself: on the Windows dev box
    Acrobat has been crashing since its 2026-08-02 update, but the printer
    driver and Distiller are separate binaries and work fine (verified
    19-Aug-2026)."""
    try:
        import win32print
        names = {p[2] for p in win32print.EnumPrinters(
            win32print.PRINTER_ENUM_LOCAL | win32print.PRINTER_ENUM_CONNECTIONS)}
        return "Adobe PDF" in names
    except Exception:
        return False


def _find_acrodist():
    """Full path of Acrobat Distiller's acrodist.exe, or None."""
    for c in (Path(r"C:/Program Files (x86)/Adobe/Acrobat DC/Acrobat/acrodist.exe"),
              Path(r"C:/Program Files/Adobe/Acrobat DC/Acrobat/acrodist.exe")):
        if c.is_file():
            return c
    return None


def _print_to_adobe_pdf(word, doc, docx_path: Path, pdf_path: Path):
    """Produce the PDF via Word -> PostScript file -> Distiller, directly.

    Word's own exporter (SaveAs FileFormat=17) cannot embed OpenType-PS (CFF)
    faces and silently rasterizes every run that uses one — with the house
    fonts that means the whole document loses its text layer, so the PDF must
    come from Adobe's pipeline instead (DOCUMENT_STANDARDS.md section 8).

    We deliberately do NOT print through the "Adobe PDF" printer *port*: its
    port monitor drops the PDF wherever its "last used folder" points, keeps
    the file locked long afterwards, and when a print collides with such a
    leftover it jams the queue with stacked modal error dialogs (19-Aug-2026
    incident). Instead, Word prints PostScript to a scratch FILE we name
    (PrintToFile), and acrodist.exe distills that file synchronously — fully
    deterministic paths, no spooler, no dialogs."""
    import subprocess
    import shutil

    acrodist = _find_acrodist()
    if acrodist is None:
        raise RuntimeError("acrodist.exe not found under Program Files — "
                           "is Acrobat/Distiller installed?")

    scratch_dir = Path(tempfile.mkdtemp(prefix="dses_distill_"))
    ps_file = scratch_dir / (pdf_path.stem + ".ps")
    out_pdf = scratch_dir / (pdf_path.stem + ".pdf")

    previous_printer = word.ActivePrinter
    try:
        word.ActivePrinter = "Adobe PDF"
        doc.PrintOut(Background=False, PrintToFile=True,
                     OutputFileName=str(ps_file))
    finally:
        try:
            word.ActivePrinter = previous_printer
        except Exception:
            pass
    # PrintOut can return before the spooler finishes writing the file
    # (observed on the install guide, 2026-08-19): poll until it exists and
    # stops growing.
    import time
    last = -1
    for _ in range(120):
        if ps_file.is_file():
            size = ps_file.stat().st_size
            if size > 0 and size == last:
                break
            last = size
        time.sleep(1.0)
    if not ps_file.is_file() or ps_file.stat().st_size == 0:
        raise RuntimeError(f"Word did not write the PostScript file {ps_file}")

    r = subprocess.run([str(acrodist), "/N", "/Q", str(ps_file)],
                       capture_output=True, text=True, timeout=300)
    if not out_pdf.is_file() or out_pdf.stat().st_size == 0:
        log = ps_file.with_suffix(".log")
        detail = log.read_text(errors="replace")[-2000:] if log.is_file() else r.stderr
        raise RuntimeError(f"Distiller produced no PDF from {ps_file}: {detail}")

    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    if pdf_path.exists():
        pdf_path.unlink()
    shutil.copy2(str(out_pdf), str(pdf_path))
    shutil.rmtree(scratch_dir, ignore_errors=True)


def _convert_docx_to_pdf_word(docx_path: Path, pdf_path: Path):
    """Open the DOCX in Word, update every field (TOC + PAGE), then write the
    PDF. Direct pywin32 instead of docx2pdf so we control the field-update
    step — without it the TOC stays as the placeholder text. Requires
    Microsoft Word on Windows.

    The PDF itself comes from the Adobe PDF printer when it is available,
    because Word's own exporter rasterizes the OpenType-PS house faces (see
    _print_to_adobe_pdf); otherwise it falls back to Word's exporter with a
    warning.

    Note: kills any leftover background winword.exe before starting. If
    you have Word open with a document you care about, save first."""
    _kill_stale_invisible_word()
    import win32com.client
    # Early binding (makepy) so NAMED arguments to PrintOut actually bind —
    # with plain dynamic Dispatch they were silently mis-delivered and the
    # PrintToFile/OutputFileName combination did nothing (19-Aug-2026).
    try:
        word = win32com.client.gencache.EnsureDispatch('Word.Application')
    except Exception:
        word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        doc = word.Documents.Open(str(docx_path.resolve()))
        try:
            # Update headers/footers + body fields. Run twice: the first
            # pass renders the TOC entries which can change page counts,
            # the second pass corrects the TOC's page numbers in light of
            # the new layout.
            for _ in range(2):
                doc.Fields.Update()
                # TablesOfContents may also need its own Update call on some
                # Word builds.
                try:
                    if doc.TablesOfContents.Count > 0:
                        doc.TablesOfContents(1).Update()
                except Exception:
                    pass
            # Word writes each TOC entry's paragraph mark in the THEME font
            # (asciiTheme=minorHAnsi, 12 pt) regardless of the toc styles, and
            # that font then rides into the PDF's embedded-font list (Cambria,
            # 2026-09-10). Pin the whole TOC range to the house heading face;
            # sizes and the level-1 bold still come from the toc styles.
            try:
                if doc.TablesOfContents.Count > 0:
                    doc.TablesOfContents(1).Range.Font.Name = FONT_HEAD
            except Exception:
                pass
            # Save the .docx so the populated TOC persists for future opens
            # in Word (otherwise the TOC reverts to the placeholder).
            doc.Save()
            if _adobe_pdf_printer_available():
                _print_to_adobe_pdf(word, doc, docx_path, pdf_path)
            else:
                print("WARNING: the 'Adobe PDF' printer was not found. Falling "
                      "back to Word's own PDF export, which RASTERIZES "
                      f"{FONT_BODY}/{FONT_HEAD} — the PDF will have no "
                      "selectable text. Install Acrobat/Distiller, or set the "
                      "FONT_* constants back to TrueType faces.",
                      file=sys.stderr)
                # wdFormatPDF = 17
                doc.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        finally:
            doc.Close(SaveChanges=False)
    finally:
        word.Quit()
    print(f"Wrote {pdf_path} ({pdf_path.stat().st_size // 1024} KB)")


def main():
    global DOC_TITLE, DOC_SUBTITLE  # cover-page/header read these as globals
    import argparse
    ap = argparse.ArgumentParser(
        description="Build a DSES-styled DOCX + PDF from a Markdown source.")
    ap.add_argument('src', nargs='?', default=str(SRC),
                    help=f"Markdown source (default: {SRC})")
    ap.add_argument('--docx', default=None,
                    help="Keep the intermediate DOCX at this path (default: a "
                         "temp file, removed after the PDF is built).")
    ap.add_argument('--pdf', default=str(DST_PDF),
                    help=f"PDF output path (default: {DST_PDF})")
    ap.add_argument('--title', default=DOC_TITLE,
                    help=f"Cover-page title (default: {DOC_TITLE!r})")
    ap.add_argument('--subtitle', default=DOC_SUBTITLE,
                    help=f"Cover-page + header subtitle "
                         f"(default: {DOC_SUBTITLE!r})")
    ap.add_argument('--version', default=DOC_VERSION, dest='doc_version',
                    help="Cover-page version line. A value starting with 'v' "
                         "renders as 'Version X.Y.Z'; anything else (e.g. "
                         f"'Rev A') is printed verbatim. Default: {DOC_VERSION!r}")
    ap.add_argument('--header-logo', default=None,
                    help="PNG shown right-justified in the page header "
                         "(pages 2+; the cover is unaffected). Default: none.")
    ap.add_argument('--no-cover', action='store_true',
                    help="Skip the cover page (short memos / one-page "
                         "instruction sheets).")
    ap.add_argument('--no-toc', action='store_true',
                    help="Skip the table of contents.")
    ap.add_argument('--toc', choices=['field', 'static'], default=None,
                    dest='toc_mode',
                    help="TOC style: 'field' = real Word TOC field with page "
                         "numbers (evaluated by Word; the default on Windows), "
                         "'static' = pre-populated outline without page numbers "
                         "(renders anywhere; the default off Windows).")
    ap.add_argument('--force', action='store_true',
                    help="Overwrite the --docx target even if git reports it "
                         "modified (i.e., discard hand-made Word edits).")
    args = ap.parse_args()
    DOC_TITLE = args.title
    DOC_SUBTITLE = args.subtitle
    globals()['DOC_VERSION'] = args.doc_version
    globals()['HEADER_LOGO'] = args.header_logo

    src = Path(args.src)
    pdf_out = Path(args.pdf)
    if not src.exists():
        print(f"Source not found: {src}", file=sys.stderr)
        sys.exit(1)

    # The DOCX is a throwaway intermediate: write it to a temp file and remove
    # it after the PDF is built, unless --docx asked to keep it somewhere.
    import tempfile
    keep_docx = args.docx is not None
    if keep_docx:
        docx_out = Path(args.docx)
    else:
        docx_out = Path(tempfile.gettempdir()) / f"{pdf_out.stem}.docx"

    # GUARD: never clobber hand-made Word edits. Rebuilding regenerates the
    # DOCX from the .md, so any direct edits to a kept, git-tracked DOCX are
    # lost. If git reports the target modified vs HEAD, someone edited it
    # since the last build -- stop and make them port the edits into the .md
    # (or pass --force to overwrite deliberately). Born of the 2026-07-13
    # trip-report incident; do not remove.
    if keep_docx and docx_out.exists() and not args.force:
        import subprocess
        try:
            r = subprocess.run(
                ['git', 'status', '--porcelain', '--', str(docx_out)],
                capture_output=True, text=True, timeout=15)
            dirty = r.returncode == 0 and r.stdout.strip().startswith(' M')
        except Exception:
            dirty = False
        if dirty:
            print(f"REFUSING to overwrite {docx_out}:\n"
                  "  git says it was modified since the last commit — it "
                  "likely contains\n  hand-made Word edits that a rebuild "
                  "would destroy.\n"
                  "  Port those edits into the Markdown source first (or "
                  "commit the DOCX),\n  then rebuild. Pass --force to "
                  "overwrite anyway.", file=sys.stderr)
            sys.exit(3)

    md_to_docx(src, docx_out, cover=not args.no_cover, toc=not args.no_toc,
               toc_mode=args.toc_mode)
    try:
        convert_docx_to_pdf(docx_out, pdf_out)
    except Exception as exc:
        print(f"PDF conversion failed: {type(exc).__name__}: {exc}",
              file=sys.stderr)
        print(f"The intermediate DOCX is at {docx_out}; open it in "
              "LibreOffice/Word and export to PDF manually.", file=sys.stderr)
        sys.exit(2)
    finally:
        if not keep_docx:
            try:
                docx_out.unlink()
            except OSError:
                pass


if __name__ == '__main__':
    main()
